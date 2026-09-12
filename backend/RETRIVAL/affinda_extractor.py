"""Affinda Resume Parser: upload file, get JSON, convert to plain text with dicts + string formatting."""

from __future__ import annotations

import mimetypes
from pathlib import Path
from typing import Any, Optional
import requests

from config import RetrievalConfig, DEFAULT_CONFIG


class AffindaResumeParser:
    """Upload document to Affinda Resume Parser → JSON → plain text."""

    API_URL = "https://resume-parser.affinda.com/v1/resumes/parse"

    def __init__(self, config: Optional[RetrievalConfig] = None, api_key: Optional[str] = None) -> None:
        cfg = config or DEFAULT_CONFIG
        self.api_key = api_key or getattr(cfg, "affinda_api_key", None) or ""

    def extract_text(self, file_path: str | Path) -> str:
        """Upload file → Affinda JSON → plain text. Falls back to local extraction if API unavailable."""
        path = Path(file_path)
        name = path.name

        if self.api_key.strip():
            try:
                print(f"[affinda] Uploading '{name}' to Affinda Resume Parser...")
                data = self._upload(path)
                text = self._json_to_text(data)
                print(f"[affinda] Converted JSON to {len(text)} chars for '{name}'.")
                print(f"[affinda] ---- text preview (first 600 chars) ----")
                for line in text[:600].splitlines():
                    print(f"[affinda]   {line}")
                print(f"[affinda] ---- end preview ----")
                return text
            except Exception as exc:
                print(f"[affinda] API error: {exc}. Falling back to local extraction.")

        print(f"[affinda] Using local text extraction for '{name}'...")
        text = self._local_extract(path)
        print(f"[affinda] Local extraction: {len(text)} chars from '{name}'.")
        return text

    # ── upload to Affinda ──────────────────────────────────────
    def _upload(self, path: Path) -> dict:
        mime, _ = mimetypes.guess_type(str(path))
        mime = mime or "application/octet-stream"

        headers = {"Authorization": f"Bearer {self.api_key.strip()}"}

        with open(path, "rb") as f:
            resp = requests.post(
                self.API_URL,
                headers=headers,
                files={"file": (path.name, f, mime)},
                timeout=120,
            )

        if resp.status_code not in (200, 201):
            raise RuntimeError(f"Affinda HTTP {resp.status_code}: {resp.text[:400]}")

        payload = resp.json()
        print(f"[affinda] Response keys: {list(payload.keys())}")
        return payload

    # ── JSON → plain text using dicts + string formatting ──────
    def _json_to_text(self, payload: dict) -> str:
        data = payload.get("data", payload)
        parts = []

        # --- Name ---
        name = self._get_val(data, "name")
        if name:
            parts.append(f"CANDIDATE NAME: {name}")

        # --- Contact ---
        contact = []
        for email in self._get_list(data, "emails"):
            contact.append(f"Email: {email}")
        for phone in self._get_list(data, "phoneNumbers"):
            contact.append(f"Phone: {phone}")
        for phone in self._get_list(data, "phones"):
            contact.append(f"Phone: {phone}")
        for link in self._get_list(data, "websites"):
            contact.append(f"Link: {link}")
        for link in self._get_list(data, "linkedin"):
            contact.append(f"LinkedIn: {link}")
        for link in self._get_list(data, "github"):
            contact.append(f"GitHub: {link}")
        loc = self._get_val(data, "location")
        if loc:
            contact.append(f"Location: {loc}")
        if contact:
            parts.append("CONTACT INFORMATION:\n" + "\n".join(contact))

        # --- Summary / Objective ---
        summary = self._get_val(data, "summary") or self._get_val(data, "objective")
        if summary:
            parts.append(f"PROFESSIONAL SUMMARY:\n{summary}")

        # --- Total experience ---
        total_exp = data.get("totalYearsExperience")
        if total_exp is not None:
            parts.append(f"TOTAL YEARS OF EXPERIENCE: {total_exp}")

        # --- Skills ---
        skills = data.get("skills", [])
        if skills:
            lines = []
            for s in skills:
                if isinstance(s, dict):
                    sname = s.get("name") or s.get("raw") or ""
                    stype = s.get("type", "")
                    months = s.get("numberOfMonths")
                    line = sname
                    if stype:
                        line += f" ({stype})"
                    if months:
                        line += f" [{months} months]"
                    if line.strip():
                        lines.append(line.strip())
                elif isinstance(s, str) and s.strip():
                    lines.append(s.strip())
            if lines:
                parts.append("SKILLS:\n" + "\n".join(lines))

        # --- Work Experience ---
        work = data.get("workExperience", []) or data.get("experience", [])
        if work:
            entries = []
            for w in work:
                if not isinstance(w, dict):
                    continue
                title = w.get("jobTitle") or w.get("rawJobTitle") or w.get("title") or "Role"
                org = w.get("organization") or w.get("employer") or w.get("company") or "Company"
                dates = w.get("dates") or {}
                start = dates.get("startDate") or w.get("startDate") or ""
                end = dates.get("endDate") or w.get("endDate") or "Present"
                desc = w.get("jobDescription") or w.get("description") or ""
                entry = f"Title: {title}\nCompany: {org}\nPeriod: {start} to {end}"
                if desc:
                    entry += f"\nResponsibilities:\n{desc}"
                entries.append(entry)
            if entries:
                parts.append("WORK EXPERIENCE:\n\n" + "\n\n".join(entries))

        # --- Education ---
        edu = data.get("education", [])
        if edu:
            entries = []
            for e in edu:
                if not isinstance(e, dict):
                    continue
                deg = e.get("accreditation") or e.get("degree") or "Degree"
                school = e.get("organization") or e.get("institution") or "Institution"
                grade = e.get("grade") or e.get("gpa") or ""
                dates = e.get("dates") or {}
                comp = dates.get("completionDate") or e.get("completionDate") or ""
                entry = f"Degree: {deg}\nInstitution: {school}"
                if grade:
                    entry += f"\nGPA/Grade: {grade}"
                if comp:
                    entry += f"\nCompleted: {comp}"
                entries.append(entry)
            if entries:
                parts.append("EDUCATION:\n\n" + "\n\n".join(entries))

        # --- Projects ---
        projects = data.get("projects", [])
        if projects:
            entries = []
            for p in projects:
                if not isinstance(p, dict):
                    continue
                title = p.get("title") or p.get("name") or "Project"
                desc = p.get("description") or ""
                entry = f"Project: {title}"
                if desc:
                    entry += f"\n{desc}"
                entries.append(entry)
            if entries:
                parts.append("PROJECTS:\n\n" + "\n\n".join(entries))

        # --- Certifications ---
        certs = data.get("certifications", [])
        if certs:
            lines = []
            for c in certs:
                if isinstance(c, dict):
                    lines.append(c.get("name") or str(c))
                elif isinstance(c, str):
                    lines.append(c)
            if lines:
                parts.append("CERTIFICATIONS:\n" + "\n".join(lines))

        # --- Achievements / Awards ---
        achievements = data.get("achievements", []) or data.get("awards", [])
        if achievements:
            lines = []
            for a in achievements:
                if isinstance(a, dict):
                    lines.append(a.get("name") or a.get("description") or str(a))
                elif isinstance(a, str):
                    lines.append(a)
            if lines:
                parts.append("ACHIEVEMENTS:\n" + "\n".join(lines))

        # --- Languages ---
        langs = data.get("languages", [])
        if langs:
            lines = []
            for l in langs:
                if isinstance(l, dict):
                    lines.append(l.get("name") or l.get("raw") or str(l))
                elif isinstance(l, str):
                    lines.append(l)
            if lines:
                parts.append("LANGUAGES:\n" + "\n".join(lines))

        # --- Raw text (full document text from Affinda) ---
        raw = data.get("rawText") or ""
        if raw.strip():
            parts.append("FULL DOCUMENT TEXT:\n" + raw.strip())

        text = "\n\n".join(parts)
        if not text.strip():
            # nothing parsed, dump whatever we got
            text = str(data)
        return text

    # ── local fallback extraction ──────────────────────────────
    def _local_extract(self, path: Path) -> str:
        suffix = path.suffix.lower()
        text = ""

        if suffix == ".pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(str(path))
                pages = [p.extract_text() or "" for p in reader.pages]
                text = "\n\n".join(p for p in pages if p.strip())
                print(f"[affinda-local] PDF: {len(text)} chars via pypdf ({len(reader.pages)} pages).")
            except Exception as e1:
                try:
                    import pypdfium2
                    pdf = pypdfium2.PdfDocument(str(path))
                    pages = [p.get_textpage().get_text_range() for p in pdf]
                    text = "\n\n".join(pages)
                    print(f"[affinda-local] PDF: {len(text)} chars via pypdfium2.")
                except Exception as e2:
                    print(f"[affinda-local] PDF failed: pypdf={e1}, pypdfium2={e2}")

        elif suffix in (".docx", ".doc"):
            try:
                import docx
                doc = docx.Document(str(path))
                paras = [p.text for p in doc.paragraphs if p.text.strip()]
                for table in doc.tables:
                    for row in table.rows:
                        cells = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                        if cells:
                            paras.append(cells)
                text = "\n".join(paras)
                print(f"[affinda-local] DOCX: {len(text)} chars.")
            except Exception as e:
                print(f"[affinda-local] DOCX failed: {e}")

        if not text.strip():
            try:
                text = path.read_text(encoding="utf-8", errors="ignore")
                print(f"[affinda-local] Plain text: {len(text)} chars.")
            except Exception:
                text = ""

        return text

    # ── helpers ─────────────────────────────────────────────────
    @staticmethod
    def _get_val(data: dict, key: str) -> str:
        """Pull a string value from a dict field that might be str or nested dict."""
        obj = data.get(key)
        if obj is None:
            return ""
        if isinstance(obj, str):
            return obj.strip()
        if isinstance(obj, dict):
            return (obj.get("raw") or obj.get("value")
                    or f"{obj.get('first', '')} {obj.get('last', '')}".strip()
                    or str(obj))
        return str(obj).strip()

    @staticmethod
    def _get_list(data: dict, key: str) -> list[str]:
        """Pull a flat list of strings from a field that might contain dicts."""
        items = data.get(key, [])
        if not isinstance(items, list):
            items = [items] if items else []
        result = []
        for item in items:
            if isinstance(item, dict):
                val = item.get("raw") or item.get("value") or item.get("url") or ""
            elif isinstance(item, str):
                val = item
            else:
                val = str(item) if item else ""
            if val and val.strip():
                result.append(val.strip())
        return result
